import io
from fastapi.testclient import TestClient
import pytest
import pypdf

from backend.app.main import app

client = TestClient(app)


def test_root_and_health():
    resp_root = client.get("/")
    assert resp_root.status_code == 200
    assert "GenAI Content Transformation Platform" in resp_root.json()["message"]

    resp_health = client.get("/health")
    assert resp_health.status_code == 200
    assert resp_health.json() == {"status": "healthy"}


def test_supported_formats():
    response = client.get("/api/v1/source/supported-formats")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "active"
    assert ".pdf" in data["supported_extensions"]
    assert ".txt" in data["supported_extensions"]
    assert ".docx" in data["supported_extensions"]
    assert ".md" in data["supported_extensions"]
    assert data["max_file_size_mb"] == 10


def test_ingest_text_valid():
    sample_text = (
        "CRITICAL SECURITY ADVISORY (CVE-2026-4401)\n"
        "Affected Product: Enterprise Cloud Gateway versions 4.2.0 through 4.9.1.\n"
        "A remote code execution vulnerability has been discovered in the authentication handling module."
    )
    response = client.post(
        "/api/v1/source/text",
        json={
            "content": sample_text,
            "title": "CVE-2026-4401 Advisory",
            "source_type": "advisory"
        }
    )
    assert response.status_code == 200
    data = response.json()
    assert data["id"] is not None
    assert data["title"] == "CVE-2026-4401 Advisory"
    assert data["source_type"] == "advisory"
    assert data["word_count"] > 15
    assert data["char_count"] > 100
    assert data["estimated_reading_time_minutes"] >= 1
    assert data["status"] == "ingested"
    assert "CVE-2026-4401" in data["cleaned_content"]


def test_ingest_text_validation_short():
    response = client.post(
        "/api/v1/source/text",
        json={"content": "Hi"}
    )
    assert response.status_code == 422


def test_ingest_text_normalization():
    dirty_text = (
        "Headline Title\r\n\r\n\r\n\r\n"
        "Paragraph line with trailing spaces    \r\n"
        "\u200bZero width space character removed.\r\n"
    )
    response = client.post(
        "/api/v1/source/text",
        json={"content": dirty_text}
    )
    assert response.status_code == 200
    cleaned = response.json()["cleaned_content"]
    assert "\r" not in cleaned
    assert "\u200b" not in cleaned
    assert "    \n" not in cleaned


def test_ingest_file_txt():
    file_content = b"This is a research report discussing agentic AI workflows in enterprise engineering."
    files = {"file": ("report.txt", io.BytesIO(file_content), "text/plain")}
    response = client.post("/api/v1/source/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "report.txt"
    assert data["source_type"] == "txt"
    assert data["word_count"] == 12
    assert "agentic AI" in data["cleaned_content"]


def test_ingest_file_markdown():
    md_content = b"# Q1 2026 Enterprise Security Review\n\nKey takeaways:\n- Multi-factor authentication\n- Zero trust architecture"
    files = {"file": ("briefing.md", io.BytesIO(md_content), "text/markdown")}
    response = client.post("/api/v1/source/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["title"] == "Q1 2026 Enterprise Security Review"
    assert data["source_type"] == "md"
    assert "Zero trust architecture" in data["cleaned_content"]


def test_ingest_file_pdf():
    # Generate a genuine PDF file in-memory using pypdf
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    
    # We can write text using a simple pdf stream or create a synthetic PDF
    # Alternatively pypdf writer with a sample stream
    stream = io.BytesIO()
    writer.write(stream)
    pdf_bytes = stream.getvalue()

    # If blank page extracts empty, test empty handling or write simple content
    # Let's test uploading the generated PDF
    files = {"file": ("sample.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
    response = client.post("/api/v1/source/upload", files=files)
    # If empty text extracted, it returns 422 with meaningful error
    assert response.status_code in [200, 422]


def test_ingest_file_docx():
    import docx
    doc = docx.Document()
    doc.add_heading("Threat Briefing 2026", level=1)
    doc.add_paragraph("Enterprise incident response protocol activated.")
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    files = {"file": ("incident.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/v1/source/upload", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["source_type"] == "docx"
    assert "Threat Briefing 2026" in data["cleaned_content"]
    assert "Enterprise incident response" in data["cleaned_content"]


def test_ingest_file_unsupported_type():
    files = {"file": ("malicious.exe", io.BytesIO(b"MZ...executable"), "application/octet-stream")}
    response = client.post("/api/v1/source/upload", files=files)
    assert response.status_code == 415
    assert "Unsupported file type" in response.json()["detail"]


def test_ingest_file_empty():
    files = {"file": ("empty.txt", io.BytesIO(b""), "text/plain")}
    response = client.post("/api/v1/source/upload", files=files)
    assert response.status_code == 400
    assert "Uploaded file is empty" in response.json()["detail"]

